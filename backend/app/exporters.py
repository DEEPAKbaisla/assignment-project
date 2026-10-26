"""DOCX export for a :class:`CustomizedResume`.

The generated document is intentionally plain: simple headings, standard
fonts, one-column layout, no tables and no images, so it is ATS-friendly.
"""
from __future__ import annotations

import io

from docx import Document
from docx.document import Document as DocumentObject
from docx.shared import Pt

from app.schemas import CustomizedResume, ExperienceItem, ProjectItem

BASE_FONT = "Calibri"
BASE_FONT_SIZE = Pt(11)


def build_docx(resume: CustomizedResume) -> bytes:
    """Render a tailored resume as a clean, ATS-friendly .docx document.

    Args:
        resume: The structured resume to export.

    Returns:
        The raw .docx bytes, ready to send as a file download.
    """
    document = Document()
    _configure_styles(document)

    _add_summary(document, resume)
    _add_skills(document, resume)
    _add_experience(document, resume)
    _add_projects(document, resume)
    _add_education(document, resume)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _configure_styles(document: DocumentObject) -> None:
    """Apply a single standard font to the document defaults."""
    normal = document.styles["Normal"]
    normal.font.name = BASE_FONT
    normal.font.size = BASE_FONT_SIZE


def _add_summary(document: DocumentObject, resume: CustomizedResume) -> None:
    """Add the professional summary section."""
    if not resume.summary.strip():
        return
    document.add_heading("Professional Summary", level=1)
    document.add_paragraph(resume.summary.strip())


def _add_skills(document: DocumentObject, resume: CustomizedResume) -> None:
    """Add the skills section as a single comma-separated line."""
    skills = [skill.strip() for skill in resume.skills if skill.strip()]
    if not skills:
        return
    document.add_heading("Skills", level=1)
    document.add_paragraph(", ".join(skills))


def _add_experience(document: DocumentObject, resume: CustomizedResume) -> None:
    """Add the work experience section with one block per role."""
    if not resume.experience:
        return
    document.add_heading("Experience", level=1)
    for item in resume.experience:
        _add_experience_item(document, item)


def _add_experience_item(document: DocumentObject, item: ExperienceItem) -> None:
    """Add a single experience entry: header line plus bullet points."""
    header = document.add_paragraph()
    header.add_run(f"{item.title} - {item.company}").bold = True
    if item.duration.strip():
        header.add_run(f" ({item.duration.strip()})")
    for bullet in item.bullets:
        if bullet.strip():
            document.add_paragraph(bullet.strip(), style="List Bullet")


def _add_projects(document: DocumentObject, resume: CustomizedResume) -> None:
    """Add the projects section with one block per project."""
    if not resume.projects:
        return
    document.add_heading("Projects", level=1)
    for project in resume.projects:
        _add_project_item(document, project)


def _add_project_item(document: DocumentObject, project: ProjectItem) -> None:
    """Add a single project entry: name plus bullet points."""
    if project.name.strip():
        document.add_paragraph(project.name.strip()).runs[0].bold = True
    for bullet in project.bullets:
        if bullet.strip():
            document.add_paragraph(bullet.strip(), style="List Bullet")


def _add_education(document: DocumentObject, resume: CustomizedResume) -> None:
    """Add the education section, one line per entry."""
    entries = [entry.strip() for entry in resume.education if entry.strip()]
    if not entries:
        return
    document.add_heading("Education", level=1)
    for entry in entries:
        document.add_paragraph(entry)
