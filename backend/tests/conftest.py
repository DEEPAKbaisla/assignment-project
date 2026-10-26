"""Shared test fixtures.

Environment variables are set *before* the app is imported so the
``GEMINI_API_KEY`` fail-fast check passes with a dummy value and tests never
call the real Gemini API.
"""
from __future__ import annotations

import os

os.environ.setdefault("GEMINI_API_KEY", "test-key-not-a-real-secret")
os.environ.setdefault("GEMINI_MODEL", "gemini-2.5-flash")
os.environ.setdefault("ALLOWED_ORIGINS", "http://localhost:3000")
os.environ.setdefault("MAX_FILE_MB", "5")
os.environ.setdefault("MIN_JD_CHARS", "100")
os.environ.setdefault("MAX_JD_CHARS", "15000")

import io

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app

VALID_JD = (
    "We are hiring a Senior Python Backend Engineer to build and maintain "
    "FastAPI microservices for our SaaS platform. You will design REST APIs, "
    "own PostgreSQL schemas, improve CI/CD pipelines, and mentor junior "
    "engineers. Must have 5+ years of Python experience, strong testing "
    "practices, and experience with Docker and AWS."
)

PAGE_ONE_LINES = [
    "Jane Doe - Senior Python Engineer",
    "Email: jane.doe@example.com | Phone: 555-0100",
    "Summary: Backend engineer with 6 years of Python experience.",
    "Skills: Python, FastAPI, Django, PostgreSQL, Docker, AWS, Git",
    "Experience:",
    "Senior Python Engineer - Acme Corp (2020 - Present)",
    "- Led migration to FastAPI, reducing p95 latency by 25%",
    "- Mentored 3 junior engineers and ran code reviews",
]

PAGE_TWO_LINES = [
    "Education:",
    "BSc Computer Science, State University, 2018",
]


def make_pdf(*pages: list[str]) -> bytes:
    """Build an in-memory PDF with one page per list of text lines."""
    document = pymupdf.open()
    for lines in pages:
        page = document.new_page()
        y_position = 72
        for line in lines:
            page.insert_text((72, y_position), line, fontsize=11)
            y_position += 16
    data = document.tobytes()
    document.close()
    return data


def make_blank_pdf() -> bytes:
    """Build a valid PDF with no text at all (like a scanned document)."""
    document = pymupdf.open()
    document.new_page()
    data = document.tobytes()
    document.close()
    return data


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    """Build an in-memory DOCX with paragraphs and an optional table."""
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client() -> TestClient:
    """FastAPI test client (server exceptions are returned as 500 JSON)."""
    return TestClient(app, raise_server_exceptions=False)


@pytest.fixture
def valid_jd() -> str:
    """A job description long enough to pass MIN_JD_CHARS."""
    return VALID_JD


@pytest.fixture
def sample_pdf() -> bytes:
    """A two-page, text-based resume PDF."""
    return make_pdf(PAGE_ONE_LINES, PAGE_TWO_LINES)


@pytest.fixture
def blank_pdf() -> bytes:
    """A valid PDF with no selectable text."""
    return make_blank_pdf()


@pytest.fixture
def sample_docx() -> bytes:
    """A DOCX resume with paragraphs, messy whitespace, and a table."""
    return make_docx(
        paragraphs=[
            "Jane Doe - Senior Python Engineer",
            "Summary:   Backend engineer with 6 years of experience. ",
            "Experience:",
            "- Led migration to FastAPI, reducing p95 latency by 25%",
        ],
        table_rows=[
            ["Skills", "Python, FastAPI, PostgreSQL"],
            ["Education", "BSc Computer Science, State University"],
        ],
    )


@pytest.fixture
def long_docx() -> bytes:
    """A DOCX with a single paragraph longer than the 20k character cap."""
    return make_docx(paragraphs=["word " * 6000])
